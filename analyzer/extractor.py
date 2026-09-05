"""
Document Text and Contact Information Extractor
Supports PDF (PyMuPDF / pdfminer / PyPDF2), DOCX, and TXT files.
"""

import re
import io
from typing import Dict, Any, Optional

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract clean text from PDF bytes using PyMuPDF (fitz) with fallbacks."""
    text = ""
    # Method 1: PyMuPDF (fitz)
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        for page in doc:
            page_text = page.get_text("text")
            if page_text:
                text += page_text + "\n"
        if text.strip():
            return clean_extracted_text(text)
    except Exception:
        pass

    # Method 2: pdfminer.six
    try:
        from pdfminer.high_level import extract_text
        text = extract_text(io.BytesIO(file_bytes))
        if text.strip():
            return clean_extracted_text(text)
    except Exception:
        pass

    # Method 3: PyPDF2 fallback
    try:
        import PyPDF2
        pdf_reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
        for page in pdf_reader.pages:
            t = page.extract_text()
            if t:
                text += t + "\n"
        if text.strip():
            return clean_extracted_text(text)
    except Exception:
        pass

    return clean_extracted_text(text)


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract clean text from DOCX bytes."""
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    paragraphs.append(" | ".join(row_text))
        return clean_extracted_text("\n".join(paragraphs))
    except Exception as e:
        return ""


def extract_text_from_file(uploaded_file) -> str:
    """Extract text based on file type."""
    if uploaded_file is None:
        return ""
    
    file_bytes = uploaded_file.getvalue() if hasattr(uploaded_file, "getvalue") else uploaded_file.read()
    filename = getattr(uploaded_file, "name", "").lower()

    if filename.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif filename.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    elif filename.endswith(".txt"):
        try:
            return clean_extracted_text(file_bytes.decode("utf-8", errors="ignore"))
        except Exception:
            return ""
    else:
        # Default try PDF first then text
        t = extract_text_from_pdf(file_bytes)
        if not t:
            try:
                t = file_bytes.decode("utf-8", errors="ignore")
            except Exception:
                pass
        return clean_extracted_text(t)


def clean_extracted_text(text: str) -> str:
    """Normalize whitespace and remove unprintable artifacts."""
    if not text:
        return ""
    # Replace weird unicode spaces/tabs
    text = re.sub(r'[\r\f\v]', '\n', text)
    text = re.sub(r'[\t\xa0]', ' ', text)
    # Remove multiple spaces on single line
    text = re.sub(r'[ ]{2,}', ' ', text)
    # Remove excessive blank lines
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()


def extract_contact_info(text: str) -> Dict[str, Any]:
    """Extract emails, phones, LinkedIn, GitHub, and portfolio URLs."""
    info = {
        "email": None,
        "phone": None,
        "linkedin": None,
        "github": None,
        "portfolio": None
    }
    
    # Email regex
    email_match = re.search(r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+', text)
    if email_match:
        info["email"] = email_match.group(0)

    # Phone regex
    phone_match = re.search(r'(\+?\d{1,3}[-.\s]?)?(\(?\d{3}\)?[-.\s]?)?\d{3}[-.\s]?\d{4}', text)
    if phone_match:
        info["phone"] = phone_match.group(0)

    # LinkedIn regex
    linkedin_match = re.search(r'(?:https?:\/\/)?(?:www\.)?linkedin\.com\/(?:in|profile)\/([a-zA-Z0-9_-]+)', text, re.IGNORECASE)
    if linkedin_match:
        info["linkedin"] = linkedin_match.group(0)

    # GitHub regex
    github_match = re.search(r'(?:https?:\/\/)?(?:www\.)?github\.com\/([a-zA-Z0-9_-]+)', text, re.IGNORECASE)
    if github_match:
        info["github"] = github_match.group(0)

    # Portfolio / Web regex
    portfolio_match = re.search(r'(?:https?:\/\/)?(?:www\.)?([a-zA-Z0-9-]+\.(?:io|me|dev|app|ai|com))(?:\/[^\s]*)?', text, re.IGNORECASE)
    if portfolio_match:
        url = portfolio_match.group(0)
        if "linkedin.com" not in url and "github.com" not in url:
            info["portfolio"] = url

    return info
